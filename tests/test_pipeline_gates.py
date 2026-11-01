r"""
Deterministic Pipeline Acceptance Gates Test Suite (tests/test_pipeline_gates.py).
Validates Gates 1 through 3:
  - Gate 1: LaTeX sanitizer escaping all TeX special chars
  - Gate 2A: 100% bijective skill proof in bullets + \entryskills{}, base language expansion, locked metrics, zero AI slop
  - Gate 2B: Single-page PDF verification (pdfinfo) and vertical pixel fill (88.0% - 98.5%)
  - Gate 3: CEO delegation email (<115 words, delegation ask), HR recruiter note, LinkedIn notes (<=300 chars)
"""

import json
import os
import re
import sys
import subprocess
import openpyxl
import pytest
from PIL import Image

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if REPO_ROOT not in sys.path:
    sys.path.insert(0, REPO_ROOT)

BANNED_AI_SLOP_WORDS = {
    "spearheaded", "synergized", "paradigm", "holistic",
    "delve", "tapestry", "utilized", "leveraged", "enthusiast"
}

FRAMEWORK_TO_BASE_LANG = {
    "fastapi": "python",
    "flask": "python",
    "langchain": "python",
    "langgraph": "python",
    "pytorch": "python",
    "scikit-learn": "python",
    "scanpy": "python",
    "react": "typescript",
    "next.js": "typescript",
    "pgvector": "sql",
    "supabase": "sql",
}

LOCKED_IMMUTABLE_METRICS = [
    "30,000+",
    "15+",
    r"89\%",
    "3.9/4.0",
]


def escape_latex_safe(plain_text: str) -> str:
    """Deterministic escaper for plain text before wrapping in LaTeX macros."""
    replacements = [
        ("\\", r"\textbackslash{}"),
        ("&", r"\&"),
        ("%", r"\%"),
        ("$", r"\$"),
        ("#", r"\#"),
        ("_", r"\_"),
        ("{", r"\{"),
        ("}", r"\}"),
        ("~", r"\textasciitilde{}"),
        ("^", r"\textasciicircum{}"),
    ]
    out = plain_text
    for old, new in replacements:
        out = out.replace(old, new)
    return out


def test_gate1_latex_escaper_handles_all_special_chars():
    raw = "Cut latency by 89% on QXO & UH_Bauer ($120K #1 tier)"
    escaped = escape_latex_safe(raw)
    assert r"89\%" in escaped
    assert r"\&" in escaped
    assert r"\_" in escaped
    assert r"\$120K" in escaped
    assert r"\#1" in escaped
    # Ensure no unescaped % remains that would comment out a LaTeX line
    assert re.search(r"(?<!\\)%", escaped) is None


def validate_applied_skills_in_tex(tex_source: str, matched_jd_skills: list[str]) -> dict:
    """Verifies that every matched JD skill is proved in bullets + \\entryskills{}."""
    tex_lower = tex_source.lower()

    # 1. Extract all \entryskills{...} blocks
    entry_skill_blocks = re.findall(r"\\entryskills\{([^}]+)\}", tex_source)
    assert len(entry_skill_blocks) >= 5, (
        f"Expected at least 5 \\entryskills{{}} blocks (roles + projects), found {len(entry_skill_blocks)}"
    )
    combined_entry_skills = " ".join(entry_skill_blocks).lower()

    # 2. Verify every matched JD skill appears in the resume body AND in \entryskills{}
    missing_from_body = [s for s in matched_jd_skills if s.lower() not in tex_lower]
    missing_from_entry_skills = [s for s in matched_jd_skills if s.lower() not in combined_entry_skills]

    # 3. Verify Base-Language Expansion
    missing_base_langs = []
    for fw, base_lang in FRAMEWORK_TO_BASE_LANG.items():
        if fw in tex_lower and base_lang not in tex_lower:
            missing_base_langs.append(f"{fw} -> missing {base_lang}")

    # 4. Verify Locked Metrics & Zero AI Slop
    missing_metrics = [m for m in LOCKED_IMMUTABLE_METRICS if m.lower() not in tex_lower]
    words = set(re.findall(r"\b[a-z]+\b", tex_lower))
    slop_found = sorted(list(words.intersection(BANNED_AI_SLOP_WORDS)))

    return {
        "entry_skills_count": len(entry_skill_blocks),
        "missing_from_body": missing_from_body,
        "missing_from_entry_skills": missing_from_entry_skills,
        "missing_base_langs": missing_base_langs,
        "missing_metrics": missing_metrics,
        "slop_found": slop_found,
        "passed": (
            not missing_from_body
            and not missing_from_entry_skills
            and not missing_base_langs
            and not missing_metrics
            and not slop_found
        ),
    }


def test_gate2a_contextual_skills_resume_passes_all_rules():
    tex_path = os.path.join(REPO_ROOT, "output", "Chandra_Rup_Daka_Resume_Contextual_Skills.tex")
    if not os.path.exists(tex_path):
        from harnesses.resume_harness import tailor_resume_for_job
        tailor_resume_for_job(1)

    tex_source = open(tex_path, encoding="utf-8").read()
    sample_jd_skills = [
        "Python", "TypeScript", "SQL", "FastAPI", "LangGraph", "LangChain",
        "AutoGen", "CrewAI", "MCP", "RAG", "Neo4j", "PostgreSQL", "Docker",
        "Google Kubernetes Engine (GKE)", "AWS Bedrock", "PyTorch"
    ]
    res = validate_applied_skills_in_tex(tex_source, sample_jd_skills)
    assert res["passed"], f"Gate 2A failed: {json.dumps(res, indent=2)}"


def test_gate2b_pdf_is_single_page_and_balanced_fill():
    pdf_path = os.path.join(REPO_ROOT, "output", "Chandra_Rup_Daka_Resume_Contextual_Skills.pdf")
    png_path = os.path.join(REPO_ROOT, "output", "Chandra_Rup_Daka_Resume_Contextual_Skills_page_1.png")

    if not os.path.exists(pdf_path) or not os.path.exists(png_path):
        from harnesses.resume_harness import tailor_resume_for_job
        tailor_resume_for_job(1)

    # 1. Check page count == 1
    info = subprocess.run(["pdfinfo", pdf_path], capture_output=True, text=True, check=True).stdout
    pages = int([line.split(":")[1].strip() for line in info.splitlines() if line.startswith("Pages:")][0])
    assert pages == 1, f"Expected strictly 1 page, got {pages}"

    # 2. Check pixel vertical fill between 88% and 98.5%
    img = Image.open(png_path).convert("RGB")
    w, h = img.size
    px = img.load()
    lowest_ink_y = 0
    for y in range(h - 1, 0, -1):
        if any(min(px[x, y]) < 245 for x in range(int(w * 0.05), int(w * 0.95), 4)):
            lowest_ink_y = y
            break
    fill_pct = round((lowest_ink_y / h) * 100.0, 1)
    assert 88.0 <= fill_pct <= 98.5, f"Vertical fill {fill_pct}% outside target [88.0%, 98.5%]"


@pytest.mark.parametrize("track,role_title,sample_skills", [
    ("agentic_swe", "Agentic Systems Engineer", [
        "Python", "TypeScript", "SQL", "FastAPI", "LangGraph", "LangChain",
        "AutoGen", "CrewAI", "MCP", "RAG", "Neo4j", "PostgreSQL", "Docker",
        "Google Kubernetes Engine (GKE)", "AWS Bedrock", "PyTorch"
    ]),
    ("core_aiml", "Machine Learning Engineer - Post-Training & Alignment", [
        "PyTorch", "Python", "SQL", "Transformers", "LoRA", "PEFT", "SFT", "DPO", "GRPO", "DeepSpeed", "FastAPI"
    ]),
    ("bioinformatics_ai", "Bioinformatics AI Engineer - Spatial Transcriptomics", [
        "Python", "Scanpy", "R", "SQL", "FastAPI", "PyTorch", "Spatial Omics", "NumPy", "SciPy"
    ]),
    ("data_analyst_ds", "Business Intelligence & AI Analytics Developer", [
        "Python", "TypeScript", "SQL", "FastAPI", "Scikit-Learn", "SHAP", "PostgreSQL", "React", "Next.Js"
    ]),
    ("forward_deployed_engineer", "Forward Deployed AI Solutions Engineer", [
        "Python", "TypeScript", "SQL", "FastAPI", "Docker", "PostgreSQL", "MCP", "Neo4j", "React"
    ]),
])
def test_gate2_across_all_five_tracks(track, role_title, sample_skills):
    """Verifies that all 5 archetype tracks compile to strictly 1 page with 88-98.5% fill and 100% bijective skill proof."""
    from harnesses.resume_harness import (
        construct_tailored_resume_ast,
        compile_resume_closed_loop,
        load_candidate_knowledge_graph
    )
    kg = load_candidate_knowledge_graph()
    job = {
        "company": f"Benchmark_{track}",
        "title": role_title,
        "assigned_track": track,
        "matched_skills_json": json.dumps(sample_skills),
        "raw_jd_text": f"Seeking {role_title} skilled in {', '.join(sample_skills)}."
    }
    ast, emap = construct_tailored_resume_ast(job, kg)
    stem = f"test_gate2_{track}"
    res = compile_resume_closed_loop(ast, stem, output_dir=os.path.join(REPO_ROOT, "output"))

    # Assertions
    assert res["page_count"] == 1, f"Expected 1 page for track {track}, got {res['page_count']}"
    assert 88.0 <= res["vertical_fill_pct"] <= 98.5, f"Fill {res['vertical_fill_pct']}% out of range for {track}"

    # Verify skills proof
    tex_content = open(res["tex_path"], encoding="utf-8").read()
    proof = validate_applied_skills_in_tex(tex_content, sample_skills)
    assert proof["passed"], f"Track {track} failed skill proof: {json.dumps(proof, indent=2)}"


def validate_outreach_draft(persona: str, email_body: str, linkedin_note: str) -> dict:
    words = email_body.strip().split()
    word_count = len(words)
    has_markdown = bool(re.search(r"(\*\*|##|`{3}|<[a-z]+>)", email_body))
    li_len = len(linkedin_note.strip())

    # CEO email must have a delegation ask ("point me to", "right person", "engineering lead", "SME")
    has_delegation_ask = True
    if persona == "CEO_DELEGATION":
        has_delegation_ask = any(
            phrase in email_body.lower()
            for phrase in ["point me to", "right person", "engineering lead", "sme", "who leads", "direct me to"]
        )

    return {
        "persona": persona,
        "word_count": word_count,
        "has_markdown": has_markdown,
        "linkedin_char_len": li_len,
        "has_delegation_ask": has_delegation_ask,
        "passed": (
            (55 <= word_count <= 125)
            and not has_markdown
            and (li_len <= 300)
            and has_delegation_ask
        ),
    }


def test_gate3_ceo_delegation_outreach_validation():
    sample_ceo_email = (
        "Hi Sarah,\n\n"
        "Saw that Ravenna is scaling out your multi-agent workflows and MCP tool integrations.\n\n"
        "Before starting my M.S. in AI at the University of Houston, I spent two years at Accenture "
        "building production Python/FastAPI multi-agent workflows (LangGraph, AutoGen) and hybrid "
        "Neo4j/pgvector RAG pipelines across 30,000+ legacy enterprise files, cutting manual impact "
        "reviews from weeks to hours. Lately I've been building live multi-provider MCP routing and "
        "LLM-as-judge evaluation harnesses across Claude, OpenAI, and Gemini.\n\n"
        "I know your plate is full, so I don't want to take up your calendar—could you point me to "
        "the engineering lead or SME owning agent infrastructure so I can share a quick note?\n\n"
        "Best,\nChandra Rup Daka\n+1 (253) 632-3181 | https://chandrarup.github.io/Portfolio/"
    )
    sample_li_note = (
        "Hi Sarah — saw Ravenna's push on multi-agent MCP workflows. I spent 2 yrs at Accenture shipping "
        "Python/FastAPI multi-agent & Neo4j/pgvector RAG systems across 30K+ files (now MS AI @ UH). "
        "Who leads agent infra on your team so I can share a quick note?"
    )
    res = validate_outreach_draft("CEO_DELEGATION", sample_ceo_email, sample_li_note)
    assert res["passed"], f"Gate 3 failed: {res}"


def test_gate3_all_database_outreach_and_excel_sync():
    """Validates that all outreach drafts in smartapply.db adhere to constraints and Excel is 100% synced."""
    from backend.db import get_db_connection, DEFAULT_DB_PATH
    excel_path = os.path.join(REPO_ROOT, "SmartApply_Master_Pipeline.xlsx")
    assert os.path.exists(excel_path), f"Master Excel file missing at: {excel_path}"

    # Verify Excel 3 tabs
    wb = openpyxl.load_workbook(excel_path)
    assert "1_Scored_Jobs_Triage" in wb.sheetnames
    assert "2_Tailored_Resumes_Audit" in wb.sheetnames
    assert "3_CEO_Manager_HR_Outreach" in wb.sheetnames
    assert wb["1_Scored_Jobs_Triage"].max_row >= 10
    assert wb["2_Tailored_Resumes_Audit"].max_row >= 5
    assert wb["3_CEO_Manager_HR_Outreach"].max_row >= 15

    # Verify DB records
    with get_db_connection(DEFAULT_DB_PATH) as conn:
        drafts = conn.execute("SELECT * FROM outreach_drafts").fetchall()
        assert len(drafts) >= 15, f"Expected at least 15 drafts in DB, found {len(drafts)}"

        for d in drafts:
            persona = d["persona_type"]
            body = d["email_body_plaintext"]
            li = d["linkedin_note_text"]
            val = validate_outreach_draft(persona, body, li)
            assert val["passed"], f"Draft ID {d['id']} failed validation: {val}"
            if persona == "CEO_DELEGATION":
                assert val["word_count"] < 115, f"CEO email exceeds 115 words ({val['word_count']} words)"

